import os
import django
from openpyxl import load_workbook

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "cna_inadimplentes.settings")
django.setup()

from nucleo.models import Devedor, Divida

arquivo = "INADIMPLENTES 2015_2025.xlsm"

wb = load_workbook(arquivo, data_only=True, keep_vba=True)
ws = wb.active

total_devedores = 0
total_dividas = 0

for row in range(3, ws.max_row + 1):
    nome = ws[f"A{row}"].value
    email = ws[f"B{row}"].value
    cpf = ws[f"C{row}"].value
    telefone = ws[f"D{row}"].value
    celular = ws[f"E{row}"].value
    parcela = ws[f"M{row}"].value
    vencimento = ws[f"N{row}"].value
    valor_original = ws[f"O{row}"].value
    valor_atual = ws[f"R{row}"].value
    categoria = ws[f"T{row}"].value

    if not nome or not cpf:
        continue

    nome = str(nome).strip()
    cpf = str(cpf).strip()
    email = str(email).strip() if email else ""
    telefone = str(celular).strip() if celular else (str(telefone).strip() if telefone else "")

    devedor, created = Devedor.objects.get_or_create(
        cpf=cpf,
        defaults={
            "nome": nome,
            "telefone": telefone,
            "email": email if email else None,
        }
    )

    if created:
        total_devedores += 1
    else:
        if not devedor.telefone and telefone:
            devedor.telefone = telefone
        if not devedor.email and email:
            devedor.email = email
        if devedor.nome != nome and nome:
            devedor.nome = nome
        devedor.save()

    ano_divida = None
    if vencimento:
        try:
            ano_divida = vencimento.year
        except:
            ano_divida = None

    Divida.objects.create(
        devedor=devedor,
        escola="CNA VIVENDAS",
        categoria_financeiro=str(categoria).strip() if categoria else "",
        ano_divida=ano_divida,
        parcela=str(parcela).strip() if parcela else "",
        vencimento=vencimento if hasattr(vencimento, "year") else None,
        valor_original=valor_original if valor_original else 0,
        valor_atual=valor_atual if valor_atual else 0,
    )

    total_dividas += 1

print(f"Importação concluída! Devedores novos: {total_devedores} | Dívidas: {total_dividas}")
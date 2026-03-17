import openpyxl
import json

file_path = "INADIMPLENTES 2015_2025.xlsm"
try:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheets = wb.sheetnames
    
    result = {
        "sheets": sheets,
        "sample_data": {}
    }
    
    for sheet_name in sheets:
        ws = wb[sheet_name]
        result["sample_data"][sheet_name] = {
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "rows": []
        }
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
            # To handle un-serializable objects (like formulas or dates if not casted well)
            row_data = []
            for cell in row:
                if cell is None:
                    row_data.append(None)
                else:
                    row_data.append(str(cell).strip())
            result["sample_data"][sheet_name]["rows"].append({"row_num": i, "values": row_data})
            
    with open("inspection_result.json", "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    print("Inspection complete. Data written to inspection_result.json")
except Exception as e:
    print("Error:", str(e))
